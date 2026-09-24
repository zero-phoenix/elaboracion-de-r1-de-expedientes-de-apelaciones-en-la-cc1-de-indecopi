"""Simulación popperiana: cada R1 del corpus es una hipótesis que el sistema debe
reproducir. Se extraen sus datos del propio texto, se regenera la R1 y se compara
párrafo a párrafo con la original. Cada diferencia es una refutación: o es error del
sistema (se corrige), o errata/variación de la original (se documenta).

  python scripts/simular.py [--solo 0057 166 …]   → pruebas/simulacion.md
"""
import argparse
import datetime as dt
import difflib
import re
import sys
import tempfile
from pathlib import Path

import docx

sys.path.insert(0, str(Path(__file__).resolve().parent))
from documento import construir as construir_docx, texto, verificar  # noqa: E402
from modelo import RAIZ, Caso, Parte, clave, fecha, fechas_en, validar_supuesto  # noqa: E402


def casar(texto_, partes):
    """Parte mencionada en un texto (por coincidencia de palabras, sin tildes)."""
    k = set(clave(texto_).split())
    mejor, pm = None, 0
    for p in partes:
        pk = set(clave(p.nombre).split()) - {"S", "A", "C", "DE", "Y", "LA", "DEL", "SEGUROS", "REASEGUROS", "COMPANIA"}
        n = len(k & pk)
        if n > pm:
            mejor, pm = p, n
    return mejor
from redaccion import GUION, PLANTILLA  # noqa: E402

TRATO = re.compile(r"^(?:a la |al |de la |del |la |el )?(señora|señor) (.+)$")


def partes_lista(s, separar_y=True):
    """«a la señora X y a Rímac …» -> [(trato, nombre)]. Solo se separa ante «y a/al/a la/el/la»
    o «e/y» seguido de señor(a), para no partir «Seguros y Reaseguros»."""
    s = s.strip()
    patron = r",? (?:y|e) (?=(?:a |al |a la |el |la |del |de la )|(?:señor|señora) )|, (?=(?:a |al |a la ))"
    trozos = re.split(patron, s) if separar_y else [s]
    out = []
    for t in trozos:
        t = t.strip().rstrip(",")
        t = re.sub(r"^(a la |al |a |de la |del |de |la |el )", lambda m: {"a la ": "la ", "al ": "el ", "de la ": "la ", "del ": "el "}.get(m.group(1), ""), t)
        m = re.match(r"^(?:la |el )?(señora|señor) (.+)$", t)
        if m:
            out.append((m.group(1), m.group(2).strip()))
        else:
            out.append(("", re.sub(r"\s+", " ", re.sub(r"^el (?=Banco)", "", t))))
    return out


def extraer(ruta):
    d = docx.Document(str(ruta))
    ps = [texto(p._p).strip() for p in d.paragraphs]
    ps = [p for p in ps if p]
    cab = {}
    for tb in d.tables:
        for r in tb.rows:
            cs = [c.text.strip() for c in r.cells]
            if len(cs) >= 3:
                cab[cs[0].split()[0].upper()] = re.sub(r"\s+", " ", cs[2])
    todo = "\n".join(ps)
    lima = next(p for p in ps if p.startswith("Lima,"))
    c = Caso(expediente=re.search(r"0*(\d+-\d{4})", cab.get("INGRESO", "")).group(1).zfill(9),
             fecha_emision=fecha(lima))
    c.expediente = f"{int(c.expediente.split('-')[0]):04d}-{c.expediente.split('-')[1]}"
    c.origen = cab.get("EXPEDIENTE", "").strip()
    m = re.search(r"(?:El |el |que el |que, el )(\d{1,2} de \w+ de \d{4}), la Comisión", todo)
    c.fecha_ingreso = fecha(m.group(1)) if m else None

    if "VISTOS" in todo and "audiencia de conciliación" in todo:
        c.supuesto = "S4"
    elif "VISTOS" in todo:
        c.supuesto = "S2"
    elif "Secretaría Técnica de la Comisión informa" in todo:
        c.supuesto = "S5"
    elif "recursos de apelación interpuestos" in todo:
        c.supuesto = "S3"
    else:
        c.supuesto = "S1"

    vias = {}
    for p in ps:
        m = re.match(r"(?:Asimismo, |Así también, |Finalmente, )?(requerir|Requerir|exhortar|Exhortar) (.+?),? para que", p)
        if not m:
            continue
        via = "casilla" if m.group(1).lower() == "exhortar" else ("domicilio" if "domicilio procesal" in p else "correo")
        for t, n in partes_lista(m.group(2)):
            vias[n] = (t, via)
    m = re.search(r"(?:interpuestos? por|presentados? por) (.+?)(?: contra la (.+?))?(?:\.|;)?$",
                  next((p for p in ps if "la Comisión de Protección al Consumidor Nº 1 ha recibido" in p), ""))
    apel_txt = m.group(1) if m else ""
    c.resolucion_apelada = (m.group(2) or "").strip() if m else ""
    ext = re.match(r"(.*?) (en el extremo .*|mediante el cual .*)$", c.resolucion_apelada)
    if ext:
        c.resolucion_apelada, c.extremo = ext.group(1), ext.group(2)
    if c.supuesto == "S5":
        m5 = re.search(r"contra la (.+?), emitida", todo)
        c.resolucion_apelada = m5.group(1) if m5 else ""
    # partes: las notificadas y las trasladadas («póngase/poner en conocimiento de X, el recurso»)
    mt = re.search(r"(?:póngase|poner) en conocimiento (?:de la |del |de )(.+?),? (?:el|los) recursos? de apelación", todo)
    if mt and not mt.group(1).startswith("las partes"):
        for n in re.split(r",\s+|\s+y (?=(?:el |la |al )?(?:Banco|señor|señora))|\s+e (?=[IH])", mt.group(1)):
            n = re.sub(r"\s+", " ", re.sub(r"^(?:el |la |al )", "", n)).strip()
            t = ""
            m2 = re.match(r"(señora|señor) (.+)", n)
            if m2:
                t, n = m2.groups()
            ya = casar(n, [Parte(nombre=x, rol="") for x in vias])
            if n and not (ya and len(set(clave(n).split()) & set(clave(ya.nombre).split())) >= 2):
                vias[n] = (t, None)
    dte_up = [x.strip() for x in cab.get("DENUNCIANTE", "").split("/")]
    for n, (t, via) in vias.items():
        juridica = re.search(r"S\.A|S\.R\.L|Banco|Asociaci|Afocat|Compañ", n)
        es_dte = bool(t) or (not juridica and any(len(set(clave(n).split()) & set(clave(u).split())) >= 2 for u in dte_up))
        c.partes.append(Parte(nombre=n, rol="dte" if es_dte else "ddo", trato=t, via=via or "correo",
                              nombre_cabecera=n.upper()))
    if c.supuesto in ("S2", "S3", "S4"):
        apel = [n for n in re.split(r",? y;? (?=(?:el |la )?(?:señor|señora|[A-ZÁÉÍÓÚ]))", apel_txt.replace(" y; ", " y "))] if apel_txt else []
    else:
        apel = [apel_txt] if apel_txt else []
    for a in apel:
        q = casar(a, c.partes)
        if q:
            q.apelante = True
    # S5: el denunciado del encabezado (no se notifica)
    if c.supuesto == "S5":
        for n in re.split(r"\s*/\s*", next((v for k, v in cab.items() if k.startswith("DENUNCIADO")), "")):
            if n:
                c.partes.append(Parte(nombre=n.title(), rol="ddo", via="correo", nombre_cabecera=n))
    for p in c.partes:
        if p.rol == "dte":
            p.nombre_cabecera = cab.get("DENUNCIANTE", p.nombre.upper())
    # escritos del apelante en S1: «el recurso de apelación y el escrito del F, presentados por»
    m = re.search(r"el recurso de apelación,? (?:y|e incluyendo posteriormente|y adicionalmente) el escrito del (\d{1,2} de \w+ de \d{4}),? presentados", todo)
    if m:
        for p in c.apelantes:
            p.escritos.append(fecha(m.group(1)))
    # VISTOS: por cláusula, la parte que la presenta y sus fechas
    mv = re.search(r"VISTOS: (.+)", todo)
    if mv:
        for cl in re.split(r";\s*(?:y\s+)?|\s+y (?=(?:el|los) escritos? del )", mv.group(1)):
            q = casar(cl, c.partes)
            if not q:
                continue
            fs = []
            for a_, b_, mes, an in re.findall(r"\b(\d{1,2}) y (\d{1,2}) de (\w+) de (\d{4})", cl):
                fs += [fecha(f"{a_} de {mes} de {an}"), fecha(f"{b_} de {mes} de {an}")]
            fs = fs or fechas_en(cl)
            if "recurso de apelación" in cl and fs:
                q.apelante = True
                q.fecha_apelacion, fs = fs[0], fs[1:]
            q.escritos += fs
    if c.supuesto == "S4":
        ma = re.search(r"se llevará a cabo el (\d+ de \w+ de \d{4}), a las ([\d:]+) horas", todo)
        if ma:
            c.audiencia_fecha, c.audiencia_hora = fecha(ma.group(1)), ma.group(2)
        ms = re.search(r"mediante el escrito del (\d+ de \w+ de \d{4}), (.+?) solicitó", todo)
        if ms:
            c.audiencia_solicitud_fecha = fecha(ms.group(1))
            c.audiencia_solicitante = ms.group(2)
    return c, ps


NORMAL = [(r"Decreto Supremo N° 004-2019-JUS", "Decreto Supremo N° 006-2026-JUS"),
          (r"LUISA ANALI SILVA MALPARTIDA|LOUSSIANA SALAZAR QUIROZ", "LOUSSIANA CATHERINE SALAZAR QUIROZ"),
          (r"Ejecutivo \d", "Especialista Legal"), (r"[A-Z]{3}/[a-zA-Z]{3}", "LSQ/DCQ"),
          (r"\s+", " "), (r"Lima, (.*?)\.?\s*$", r"Lima, \1")]


def norm(p):
    for a, b in NORMAL:
        p = re.sub(a, b, p)
    return p.strip()


def comparar(orig, gen):
    o = [norm(p) for p in orig if norm(p)]
    g = [norm(p) for p in gen if norm(p)]
    ratio = difflib.SequenceMatcher(None, "\n".join(o), "\n".join(g)).ratio()
    dif = [l for l in difflib.unified_diff(o, g, "original", "generada", n=0, lineterm="") if l[:1] in "+-" and l[:3] not in ("+++", "---")]
    return ratio, dif


def main():
    a = argparse.ArgumentParser()
    a.add_argument("--solo", nargs="*")
    a.add_argument("--salida", default=str(RAIZ / "pruebas" / "simulacion.md"))
    a.add_argument("--guardar", default=None, help="carpeta donde dejar las R1 regeneradas")
    A = a.parse_args()
    fs = sorted((RAIZ / "corpus" / "r1").glob("*.docx"))
    L = ["# Simulación popperiana sobre el corpus", "",
         "Cada R1 revisada del corpus se reproduce con el sistema a partir de sus propios datos. "
         "Se normalizan solo las diferencias ya decididas (firma vigente, DS 006-2026-JUS, iniciales, punto tras la fecha).", "",
         "| R1 | Supuesto | Similitud | Verificador | Diferencias |", "|---|---|---|---|---|"]
    detalle = []
    tot = []
    for f in fs:
        if A.solo and not any(s in f.name for s in A.solo):
            continue
        try:
            c, orig = extraer(f)
            validar_supuesto(c)
            if c.observaciones:
                L.append(f"| {f.name} | {c.supuesto} | — | OBSERVADO por el validador (datos extraídos incompletos): {'; '.join(c.observaciones)} | — |")
                continue
            with tempfile.TemporaryDirectory() as td:
                out = Path(td) / "g.docx"
                construir_docx(c, GUION[c.supuesto](c), RAIZ / "plantillas" / PLANTILLA[c.supuesto], out)
                err = verificar(out, c)
                if A.guardar:
                    Path(A.guardar).mkdir(parents=True, exist_ok=True)
                    import shutil
                    shutil.copy(out, Path(A.guardar) / f"SIM {f.name}")
                g = [texto(p._p).strip() for p in docx.Document(str(out)).paragraphs]
            r, dif = comparar(orig, g)
            tot.append(r)
            L.append(f"| {f.name} | {c.supuesto} | {r:.3f} | {'APTO' if not err else 'NO APTO: ' + '; '.join(err)} | {len(dif)} |")
            if dif:
                detalle += [f"### {f.name} ({c.supuesto}, {r:.3f})", "```diff"] + [l[:400] for l in dif] + ["```", ""]
        except Exception as e:  # una R1 que no se puede extraer también es una refutación
            L.append(f"| {f.name} | ? | — | ERROR: {type(e).__name__}: {e} | — |")
    if tot:
        L.insert(4, f"Media de similitud: **{sum(tot) / len(tot):.3f}** sobre {len(tot)} R1 reproducidas.\n")
    Path(A.salida).write_text("\n".join(L + ["", "## Diferencias", ""] + detalle) + "\n", encoding="utf-8")
    print("\n".join(L[:6 + len(fs)]))


if __name__ == "__main__":
    main()
