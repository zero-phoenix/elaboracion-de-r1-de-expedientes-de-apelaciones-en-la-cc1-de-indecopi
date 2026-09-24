"""Datos de un caso de R1 de apelación: lectura del Excel de control
(hoja «Recibidas 2026») y de la hoja complementaria R1_DATOS, validación y
deducción del supuesto. No redacta: solo reúne y comprueba datos."""
import datetime as dt
import json
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
MESES = {"enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6, "julio": 7,
         "agosto": 8, "setiembre": 9, "septiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12}
VIAS = {"correo", "casilla", "domicilio"}
SUPUESTOS = {
    "S1": "Simple: una apelación (con o sin escritos del propio apelante)",
    "S2": "Apelación(es) y escritos de la otra parte (VISTOS / CONSIDERANDO / SE RESUELVE)",
    "S3": "Dos apelaciones, sin escritos",
    "S4": "Con audiencia de conciliación",
    "S5": "Apelación contra inadmisibilidad o improcedencia liminar: solo se notifica al apelante",
}


# ------------------------------------------------------------ utilidades ---
def clave(s):
    s = unicodedata.normalize("NFD", str(s or "").upper())
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", re.sub(r"[^A-Z0-9 ]", " ", s)).strip()


def limpio(v):
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v).replace("\xa0", " ")).strip()


def num_exp(s):
    """«000663-2026/CC1-APELACION» / «663-2026» -> «0663-2026»."""
    m = re.search(r"0*(\d+)-(\d{4})", str(s or ""))
    return f"{int(m.group(1)):04d}-{m.group(2)}" if m else None


def fecha(v):
    """date/datetime/«12/06/2026»/«12-06-2026»/«9 de febrero de(l) 2026» -> date (o None)."""
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = limpio(v).lower()
    m = re.search(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})", s)
    if m:
        return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    m = re.search(r"(\d{1,2}) de (\w+) del? (\d{4})", s)
    if m and m.group(2) in MESES:
        return dt.date(int(m.group(3)), MESES[m.group(2)], int(m.group(1)))
    return None


def fechas_en(texto):
    """Todas las fechas de un texto libre, en orden de aparición."""
    out = []
    s = limpio(texto).lower()
    for m in re.finditer(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})|(\d{1,2}) de (\w+) del? (\d{4})", s):
        if m.group(1):
            out.append(dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1))))
        elif m.group(5) in MESES:
            out.append(dt.date(int(m.group(6)), MESES[m.group(5)], int(m.group(4))))
    return out


def nombre_propio(s):
    """«HERNAN RODRIGO MAR PEREZ» -> «Hernan Rodrigo Mar Perez» (no añade tildes)."""
    return " ".join(w.capitalize() if not re.fullmatch(r"(DE|DEL|LA|LAS|LOS|Y)", w) else w.lower()
                    for w in limpio(s).split())


def directorio():
    return json.loads((RAIZ / "docs" / "directorio_notificacion.json").read_text(encoding="utf-8"))["proveedores"]


def buscar_proveedor(nombre):
    """Busca en el directorio por clave contenida en el nombre (Excel o texto)."""
    k = clave(nombre).replace("PERU", "").replace("  ", " ")
    mejor = None
    for clv, dat in directorio().items():
        c = clave(clv).replace("PERU", "").replace("  ", " ").strip()
        if c and c in k and (mejor is None or len(c) > len(mejor[0])):
            mejor = (c, clv, dat)
    return (mejor[1], mejor[2]) if mejor else (None, None)


# ------------------------------------------------------------------ datos ---
@dataclass
class Parte:
    nombre: str                  # como va en el texto: «Hernan Rodrigo Mar Perez», «Rímac Seguros y Reaseguros S.A.»
    rol: str                     # «dte» | «ddo»
    trato: str = ""              # «señor» | «señora» | «» (persona jurídica)
    via: str = ""                # correo | casilla | domicilio
    apelante: bool = False
    fecha_apelacion: dt.date = None
    escritos: list = field(default_factory=list)   # fechas de escritos a trasladar
    nombre_cabecera: str = ""    # MAYÚSCULAS para la tabla de encabezado


@dataclass
class Caso:
    expediente: str              # «0663-2026»
    origen: str = ""             # «0010-2026/PS1»
    fecha_ingreso: dt.date = None
    partes: list = field(default_factory=list)
    resolucion_apelada: str = ""  # «Resolución Final Nº 2157-2025/PS1»
    extremo: str = ""             # «en el extremo que …» (opcional)
    supuesto: str = ""
    audiencia_fecha: dt.date = None
    audiencia_hora: str = ""
    audiencia_solicitante: str = ""   # «dte» | «ddo» | nombre
    audiencia_solicitud_fecha: dt.date = None
    fecha_emision: dt.date = None
    observaciones: list = field(default_factory=list)   # bloquean
    advertencias: list = field(default_factory=list)    # no bloquean; se informan

    @property
    def denunciantes(self):
        return [p for p in self.partes if p.rol == "dte"]

    @property
    def denunciados(self):
        return [p for p in self.partes if p.rol == "ddo"]

    @property
    def apelantes(self):
        return [p for p in self.partes if p.apelante]

    @property
    def no_apelantes(self):
        return [p for p in self.partes if not p.apelante]

    @property
    def expediente_cc1(self):
        return f"{self.expediente}/CC1-APELACIÓN"


# ------------------------------------------------------------ lectura ---
def _hoja(wb, nombre):
    for ws in wb.worksheets:
        if clave(ws.title) == clave(nombre):
            return ws
    raise SystemExit(f"No encuentro la hoja «{nombre}» en el Excel. Hojas: {wb.sheetnames}")


def leer_tabla(ruta, hoja):
    """Filas como dict {ENCABEZADO_NORMALIZADO: valor}; ubica columnas por nombre."""
    wb = openpyxl.load_workbook(ruta, data_only=True)
    ws = _hoja(wb, hoja)
    filas = list(ws.iter_rows(values_only=True))
    ih = next(i for i, f in enumerate(filas) if any(clave(c) in ("N DE EXPEDIENTE", "EXPEDIENTE") for c in f))
    enc = [clave(c) for c in filas[ih]]
    out = []
    for f in filas[ih + 1:]:
        d = {}
        for i, c in enumerate(enc):
            if c and i < len(f):
                d.setdefault(c, f[i])
        out.append(d)
    return out


def control(ruta_excel, hoja="Recibidas 2026"):
    out = {}
    for d in leer_tabla(ruta_excel, hoja):
        n = num_exp(d.get("N DE EXPEDIENTE"))
        if n:
            out[n] = d
    return out


def datos_r1(ruta):
    out = {}
    if not Path(ruta).exists():
        return out
    for d in leer_tabla(ruta, "R1_DATOS"):
        n = num_exp(d.get("EXPEDIENTE"))
        if n:
            out[n] = {k.replace(" ", "_"): v for k, v in d.items() if limpio(v) != ""}
    return out


def _partir(v):
    return [limpio(x) for x in re.split(r"\s*/\s*", limpio(v)) if limpio(x) and limpio(x) != "-"]


# -------------------------------------------------------- construcción ---
def construir(exp, fila, extra, fecha_emision):
    """Une control + R1_DATOS en un Caso; toda falta queda en observaciones."""
    c = Caso(expediente=exp, fecha_emision=fecha_emision)
    ob, adv = c.observaciones, c.advertencias
    X = lambda k: extra.get(k)  # noqa: E731

    c.origen = limpio(X("EXPEDIENTE_ORIGEN") or (fila or {}).get("EXP ORIGEN"))
    if not re.fullmatch(r"\d{1,5}-\d{4}/PS\d", c.origen or ""):
        ob.append(f"Expediente de origen ausente o mal escrito: «{c.origen}».")

    c.fecha_ingreso = fecha(X("FECHA_INGRESO_CC1"))
    if not c.fecha_ingreso:
        ob.append("Falta FECHA_INGRESO_CC1: la fecha en que la CC1 recibió el expediente "
                  "(no coincide con las columnas de fecha del Excel de control).")

    # --- denunciantes
    dte_nom = _partir(X("DENUNCIANTE_NOMBRE")) or [nombre_propio(x) for x in _partir((fila or {}).get("DENUNCIANTE S"))]
    if not X("DENUNCIANTE_NOMBRE") and dte_nom:
        adv.append("Nombre del denunciante tomado del Excel en mayúsculas: revisar tildes (DENUNCIANTE_NOMBRE).")
    tratos = _partir(X("DENUNCIANTE_TRATO"))
    vias = [v.lower() for v in _partir(X("DENUNCIANTE_VIA"))]
    for i, n in enumerate(dte_nom):
        t = (tratos[i] if i < len(tratos) else "").lower()
        if t not in ("señor", "señora", "juridica", "jurídica"):
            ob.append(f"Falta DENUNCIANTE_TRATO (señor/señora) para «{n}».")
        v = vias[i] if i < len(vias) else ""
        if not v:
            v = "correo"
            adv.append(f"Vía del denunciante «{n}» no indicada: se usa correo (la de 47 de 49 denunciantes del corpus). "
                       "Si la cédula dice CASILLA-E o domicilio, indicarlo en DENUNCIANTE_VIA.")
        c.partes.append(Parte(nombre=n, rol="dte", trato=t if t in ("señor", "señora") else "", via=v,
                              nombre_cabecera=n.upper()))
    if not dte_nom:
        ob.append("Sin denunciante.")

    # --- denunciados
    ddo_src = _partir(X("DENUNCIADOS_NOMBRES")) or _partir((fila or {}).get("DENUNCIADO S"))
    ddo_vias = [v.lower() for v in _partir(X("DENUNCIADOS_VIAS"))]
    for i, n in enumerate(ddo_src):
        k, dat = buscar_proveedor(n)
        nombre = dat["razon_social"] if dat and not X("DENUNCIADOS_NOMBRES") else (n if X("DENUNCIADOS_NOMBRES") else nombre_propio(n))
        v = ddo_vias[i] if i < len(ddo_vias) else ""
        if not v:
            if dat and dat.get("via") and not dat.get("mixta"):
                v = dat["via"]
            elif dat and dat.get("mixta"):
                ob.append(f"«{nombre}» tiene vía MIXTA en el corpus (correo y casilla): indicar DENUNCIADOS_VIAS según la cédula/expediente.")
            else:
                ob.append(f"«{nombre}» no figura en el directorio de notificación: indicar DENUNCIADOS_VIAS (correo/casilla/domicilio).")
        if not dat and not X("DENUNCIADOS_NOMBRES"):
            ob.append(f"Razón social de «{n}» no está en el directorio: escribirla exacta en DENUNCIADOS_NOMBRES.")
        c.partes.append(Parte(nombre=nombre, rol="ddo", via=v, nombre_cabecera=nombre.upper()))
    if not ddo_src:
        ob.append("Sin denunciado.")
    for p in c.partes:
        if p.via and p.via not in VIAS:
            ob.append(f"Vía no válida para «{p.nombre}»: «{p.via}» (correo/casilla/domicilio).")

    # --- apelantes
    ap_txt = _partir(X("APELANTES")) or _partir((fila or {}).get("APELANTE S"))
    f_ap = [fecha(x) for x in _partir(X("FECHAS_APELACION"))]
    esc_libre = limpio((fila or {}).get("ESCRITOS A TRASLADAR"))
    if not ap_txt:
        ob.append("No consta quién apela (APELANTES: «DTE», «DDO», «DDO2» o el nombre).")
    for i, a in enumerate(ap_txt):
        p = _identificar(c, a)
        if p is None:
            ob.append(f"El apelante «{a}» no es parte según el encabezado (denunciante/denunciados): consultar.")
            continue
        p.apelante = True
        if i < len(f_ap) and f_ap[i]:
            p.fecha_apelacion = f_ap[i]

    # escritos: R1_DATOS manda; si no, se intentan leer de la columna libre del Excel (se advierte)
    if X("ESCRITOS") is not None or X("FECHAS_APELACION") is not None:
        for seg in re.split(r"\s*;\s*", limpio(X("ESCRITOS") or "")):
            if not seg:
                continue
            m = re.match(r"(\S+)\s+(.*)", seg)
            p = _identificar(c, m.group(1)) if m else None
            fs = fechas_en(m.group(2)) if m else []
            if p is None or not fs:
                ob.append(f"ESCRITOS mal escrito: «{seg}» (formato: «DTE 12/06/2026; DDO 19/12/2025»).")
                continue
            p.escritos.extend(fs)
    elif esc_libre:
        _escritos_desde_excel(c, esc_libre)
        adv.append(f"Fechas de apelación/escritos leídas de «ESCRITOS A TRASLADAR»: «{esc_libre}». Confirmar.")

    c.resolucion_apelada = limpio(X("RESOLUCION_APELADA"))
    c.extremo = limpio(X("EXTREMO"))
    c.audiencia_fecha = fecha(X("AUDIENCIA_FECHA"))
    c.audiencia_hora = limpio(X("AUDIENCIA_HORA"))
    sol = limpio(X("AUDIENCIA_SOLICITUD"))
    if sol:
        m = re.match(r"(\S+)\s+(.*)", sol)
        c.audiencia_solicitante = m.group(1) if m else sol
        c.audiencia_solicitud_fecha = fecha(m.group(2)) if m else None

    if limpio(X("OBSERVACIONES")).upper().startswith("PENDIENTE"):
        ob.append("Marcado PENDIENTE en R1_DATOS: " + limpio(X("OBSERVACIONES")))
    ya = set()
    for p in c.apelantes:
        if id(p) in ya:
            ob.append(f"«{p.nombre}» figura dos veces como apelante: usar DTE / DTE2 o DDO / DDO2.")
        ya.add(id(p))
    c.supuesto = limpio(X("SUPUESTO")).upper() or deducir_supuesto(c)
    if c.supuesto not in SUPUESTOS:
        ob.append(f"Supuesto desconocido: «{c.supuesto}».")
    validar_supuesto(c)
    return c


def _identificar(c, texto):
    """«DTE», «DDO», «DDO2», nombre (o parte del nombre) -> Parte."""
    k = clave(texto)
    m = re.fullmatch(r"(?:DTE|DENUNCIANTE)(\d?)", k)
    if m:
        i = int(m.group(1) or 1) - 1
        return c.denunciantes[i] if i < len(c.denunciantes) else None
    m = re.fullmatch(r"DDO(\d?)", k)
    if m:
        i = int(m.group(1) or 1) - 1
        return c.denunciados[i] if i < len(c.denunciados) else None
    kset = set(k.split())
    for p in c.partes:
        pk = set(clave(p.nombre).split())
        if kset and (kset <= pk or (len(kset & pk) >= 2)):
            return p
    for p in c.denunciados:
        kp, _ = buscar_proveedor(p.nombre)
        k2, _ = buscar_proveedor(texto)
        if kp and kp == k2:
            return p
    return None


def _escritos_desde_excel(c, texto):
    """Lectura prudente de la columna libre «ESCRITOS A TRASLADAR»."""
    for seg in re.split(r"\s*;\s*|\s+-\s+", texto):
        fs = fechas_en(seg)
        if not fs:
            continue
        s = clave(seg)
        quien = None
        if re.search(r"\bDTE\b|DENUNCIANTE", s):
            quien = c.denunciantes[0] if c.denunciantes else None
        elif re.search(r"\bDDO\b|\bDO\b", s):
            quien = c.denunciados[0] if c.denunciados else None
        else:
            for p in c.partes:
                if any(w in s.split() for w in clave(p.nombre).split() if len(w) > 4):
                    quien = p
                    break
            if quien is None and "APELANTE" in s and len(c.apelantes) == 1:
                quien = c.apelantes[0]
        if quien is None:
            c.observaciones.append(f"No sé de quién es el escrito «{seg}»: indicarlo en ESCRITOS.")
            continue
        if "APELACI" in s and quien.fecha_apelacion is None:
            quien.fecha_apelacion = fs[0]
        else:
            quien.escritos.append(fs[0])


def deducir_supuesto(c):
    if c.audiencia_fecha:
        return "S4"
    otros_escritos = any(p.escritos for p in c.no_apelantes)
    if otros_escritos or (len(c.apelantes) > 1 and any(p.escritos for p in c.partes)):
        return "S2"
    if len({p.rol for p in c.apelantes}) > 1:
        return "S3"   # recursos de partes distintas
    return "S1"       # un recurso (también el conjunto de varios denunciantes)


def validar_supuesto(c):
    ob = c.observaciones
    s = c.supuesto
    if s in ("S2", "S3", "S4", "S5") and not c.resolucion_apelada:
        ob.append(f"{s} exige RESOLUCION_APELADA (p. ej. «Resolución Final Nº 2157-2025/PS1»).")
    if s == "S3" and len(c.apelantes) < 2:
        ob.append("S3 exige dos apelantes.")
    if s in ("S2", "S4") and any(p.apelante and not p.fecha_apelacion for p in c.partes):
        ob.append(f"{s} cita la fecha de cada recurso en VISTOS: falta FECHAS_APELACION.")
    if s == "S4":
        if not (c.audiencia_fecha and c.audiencia_hora):
            ob.append("S4 exige AUDIENCIA_FECHA y AUDIENCIA_HORA.")
        if not (c.audiencia_solicitante and c.audiencia_solicitud_fecha):
            ob.append("S4 exige AUDIENCIA_SOLICITUD («DDO 08/01/2026»: quién la pidió y la fecha de su escrito).")
    if s == "S5":
        if len(c.apelantes) != 1 or c.apelantes[0].rol != "dte":
            ob.append("S5: el único apelante es el denunciante (no hubo admisión ni imputación).")
    for p in c.partes:
        if s != "S5" or p.apelante:
            if not p.via:
                pass  # ya observado
    if s in ("S1", "S2", "S4") and not c.no_apelantes and len(c.apelantes) < 2:
        ob.append("No hay parte no apelante a quien trasladar el recurso.")
    if s == "S1" and len({p.rol for p in c.apelantes}) > 1:
        ob.append("S1 es para un solo recurso: apelan partes distintas (S3, o S2 si hay escritos).")
