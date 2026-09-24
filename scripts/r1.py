"""R1 de expedientes de apelación — CC1 Indecopi.

Uso (desde la raíz del repositorio):
  python scripts/r1.py datos    --exps 0663-2026 0685-2026 …   crea/completa entrada/R1_DATOS.xlsx
  python scripts/r1.py leer     [--exps …]                       salida/_REPORTE.md (APTOS / OBSERVADOS)
  python scripts/r1.py generar  [--exps …] [--fecha DD/MM/AAAA]  salida/R1 NNNN-2026-CC1-APE.docx (+ verificación)
  python scripts/r1.py verificar <docx> …                        APTO / NO APTO
Opciones comunes: --excel (entrada/APES_R1_SEGUROS.xlsx) --hoja («Recibidas 2026») --datos (entrada/R1_DATOS.xlsx)
"""
import argparse
import json
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent))
from documento import construir as construir_docx  # noqa: E402
from documento import verificar  # noqa: E402
from modelo import (RAIZ, SUPUESTOS, construir, control, datos_r1, fecha, limpio,  # noqa: E402
                    num_exp)
from redaccion import GUION, PLANTILLA  # noqa: E402

COLUMNAS = [
    ("EXPEDIENTE", "número de ingreso NNNN-AAAA"),
    ("SUPUESTO", "S1…S5 (vacío = se deduce)"),
    ("FECHA_EMISION", "DD/MM/AAAA solo si este expediente lleva fecha propia (vacío = la de la remesa)"),
    ("FECHA_INGRESO_CC1", "DD/MM/AAAA: fecha en que la CC1 recibió el expediente"),
    ("EXPEDIENTE_ORIGEN", "vacío = el del Excel de control"),
    ("DENUNCIANTE_NOMBRE", "Como va en el texto, con tildes (varios: /)"),
    ("DENUNCIANTE_TRATO", "señor / señora (varios: /)"),
    ("DENUNCIANTE_VIA", "correo / casilla / domicilio (vacío = correo)"),
    ("DENUNCIADOS_NOMBRES", "vacío = razón social del directorio"),
    ("DENUNCIADOS_VIAS", "correo / casilla / domicilio, en el orden de los denunciados (vacío = directorio)"),
    ("APELANTES", "DTE / DDO / DDO2 (dos: «DTE / DDO»)"),
    ("FECHAS_APELACION", "DD/MM/AAAA, en el orden de APELANTES"),
    ("RESOLUCION_APELADA", "«Resolución Final Nº 2157-2025/PS1» (S5: «Resolución 3 de fecha 11 de julio de 2025»)"),
    ("EXTREMO", "opcional: «en el extremo que …»"),
    ("ESCRITOS", "escritos a trasladar distintos del recurso: «DTE 25/06/2026; DDO 19/12/2025»"),
    ("AUDIENCIA_FECHA", "S4: DD/MM/AAAA"),
    ("AUDIENCIA_HORA", "S4: «16:00»"),
    ("AUDIENCIA_SOLICITUD", "S4: «DDO 08/01/2026» (quién la pidió y fecha del escrito)"),
    ("OBSERVACIONES", "libre"),
]


def args():
    a = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    a.add_argument("accion", choices=["datos", "leer", "generar", "verificar"])
    a.add_argument("docx", nargs="*")
    a.add_argument("--excel", default=str(RAIZ / "entrada" / "APES_R1_SEGUROS.xlsx"))
    a.add_argument("--hoja", default="Recibidas 2026")
    a.add_argument("--datos", default=str(RAIZ / "entrada" / "R1_DATOS.xlsx"))
    a.add_argument("--exps", nargs="*", default=None)
    a.add_argument("--fecha", default=None)
    a.add_argument("--salida", default=str(RAIZ / "salida"))
    a.add_argument("--sin-vista", action="store_true", help="no generar capturas ni fichas de formato")
    return a.parse_args()


def fecha_emision(a):
    if a.fecha:
        return fecha(a.fecha)
    v = json.loads((RAIZ / "config" / "remesa.json").read_text(encoding="utf-8")).get("fecha_emision")
    return fecha(v) if v else None


def casos(a):
    ctrl = control(a.excel, a.hoja)
    extra = datos_r1(a.datos)
    exps = [num_exp(e) for e in a.exps] if a.exps else sorted(extra)
    fe = fecha_emision(a)
    out = []
    for e in exps:
        c = construir(e, ctrl.get(e), extra.get(e, {}), fe)
        if e not in ctrl:
            c.advertencias.append("No figura en la hoja de control del Excel.")
        if e not in extra:
            c.observaciones.insert(0, "Sin fila en R1_DATOS.")
        out.append(c)
    return out, fe


def cmd_datos(a):
    ruta = Path(a.datos)
    if ruta.exists():
        wb = openpyxl.load_workbook(ruta)
        ws = wb["R1_DATOS"]
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "R1_DATOS"
        ws.append([c for c, _ in COLUMNAS])
        ws.append([d for _, d in COLUMNAS])
        for c in ws[1]:
            c.font = Font(bold=True)
        for c in ws[2]:
            c.font = Font(italic=True, color="808080")
            c.fill = PatternFill("solid", fgColor="F2F2F2")
    ya = {num_exp(r[0]) for r in ws.iter_rows(min_row=3, values_only=True) if r and r[0]}
    ctrl = control(a.excel, a.hoja)
    n = 0
    for e in [num_exp(x) for x in (a.exps or [])]:
        if e in ya:
            continue
        fila = ctrl.get(e, {})
        ws.append(fila_datos(e, fila))
        n += 1
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)
    print(f"{ruta}: {n} filas nuevas.")


def _r1_corpus(e):
    """R1 ya revisada del corpus para ese expediente (control), si existe."""
    num = int(e.split("-")[0])
    for f in sorted((RAIZ / "corpus" / "r1").glob("*.docx")):
        m = re.search(r"(\d{2,4})-2026", f.name.replace("1 608", "608"))
        if m and int(m.group(1)) == num and e.endswith("2026"):
            return f
    return None


def etiqueta(c, p):
    """DTE, DTE2…, DDO, DDO2… según su orden en el encabezado."""
    grupo = c.denunciantes if p.rol == "dte" else c.denunciados
    i = grupo.index(p)
    return ("DTE" if p.rol == "dte" else "DDO") + (str(i + 1) if i else "")


def fila_datos(e, fila):
    """Fila de R1_DATOS: desde la R1 del corpus si existe (control); si no, solo lo que el Excel afirma."""
    from modelo import Caso, _escritos_desde_excel, construir as _c
    cols = {c: "" for c, _ in COLUMNAS}
    cols["EXPEDIENTE"] = e
    f = _r1_corpus(e)
    if f is not None:
        from simular import extraer
        c, _ = extraer(f)
        cols.update({
            "SUPUESTO": c.supuesto, "FECHA_INGRESO_CC1": c.fecha_ingreso.strftime("%d/%m/%Y") if c.fecha_ingreso else "",
            "EXPEDIENTE_ORIGEN": c.origen,
            "DENUNCIANTE_NOMBRE": " / ".join(p.nombre for p in c.denunciantes),
            "DENUNCIANTE_TRATO": " / ".join(p.trato for p in c.denunciantes),
            "DENUNCIANTE_VIA": " / ".join(p.via for p in c.denunciantes),
            "DENUNCIADOS_NOMBRES": " / ".join(p.nombre for p in c.denunciados),
            "DENUNCIADOS_VIAS": " / ".join(p.via for p in c.denunciados),
            "APELANTES": " / ".join(etiqueta(c, p) for p in c.apelantes),
            "FECHAS_APELACION": " / ".join(p.fecha_apelacion.strftime("%d/%m/%Y") for p in c.apelantes if p.fecha_apelacion),
            "RESOLUCION_APELADA": c.resolucion_apelada, "EXTREMO": c.extremo,
            "ESCRITOS": "; ".join(f"{etiqueta(c, p)} {d.strftime('%d/%m/%Y')}" for p in c.partes for d in p.escritos),
            "OBSERVACIONES": f"CONTROL: datos tomados de la R1 ya revisada «{f.name}» (corpus/r1).",
        })
        return [cols[c] for c, _ in COLUMNAS]
    if not fila:
        cols["OBSERVACIONES"] = "No figura en el Excel de control."
        return [cols[c] for c, _ in COLUMNAS]
    cols["EXPEDIENTE_ORIGEN"] = limpio(fila.get("EXP ORIGEN"))
    c = _c(e, fila, {}, None)
    ap = [p for p in c.apelantes]
    if ap:
        cols["APELANTES"] = " / ".join(etiqueta(c, p) for p in ap)
        cols["FECHAS_APELACION"] = " / ".join(p.fecha_apelacion.strftime("%d/%m/%Y") for p in ap if p.fecha_apelacion)
    cols["ESCRITOS"] = "; ".join(f"{etiqueta(c, p)} {d.strftime('%d/%m/%Y')}" for p in c.partes for d in p.escritos)
    cols["OBSERVACIONES"] = "Prellenado solo con lo que afirma el Excel de control; completar las columnas vacías."
    return [cols[c] for c, _ in COLUMNAS]


def reporte(cs, fe, generados=None):
    generados = generados or {}
    L = ["# Reporte de R1 de apelaciones", "",
         f"Fecha de emisión: {fe.strftime('%d/%m/%Y') if fe else 'SIN FIJAR (config/remesa.json o --fecha)'}", "",
         "| Expediente | Supuesto | Estado | Detalle |", "|---|---|---|---|"]
    for c in cs:
        estado = "OBSERVADO" if c.observaciones else "APTO"
        if c.expediente in generados:
            estado = generados[c.expediente]
        det = "<br>".join(["**" + o + "**" for o in c.observaciones] + c.advertencias) or "—"
        L.append(f"| {c.expediente} | {c.supuesto} | {estado} | {det} |")
    L += ["", "Supuestos: " + "; ".join(f"{k} = {v}" for k, v in SUPUESTOS.items())]
    return "\n".join(L) + "\n"


def cmd_leer(a, generados=None):
    cs, fe = casos(a)
    out = Path(a.salida)
    out.mkdir(exist_ok=True)
    (out / "_REPORTE.md").write_text(reporte(cs, fe, generados), encoding="utf-8")
    for c in cs:
        print(f"{c.expediente}  {c.supuesto:3}  {'OBSERVADO' if c.observaciones else 'APTO'}")
        for o in c.observaciones:
            print(f"    ✗ {o}")
        for w in c.advertencias:
            print(f"    · {w}")
    print(f"Reporte: {out / '_REPORTE.md'}")
    return cs, fe


def cmd_generar(a):
    cs, fe = casos(a)
    for c in cs:
        if c.fecha_emision is None:
            c.observaciones.append("Fecha de emisión sin fijar (FECHA_EMISION, config/remesa.json o --fecha).")
    out = Path(a.salida)
    out.mkdir(exist_ok=True)
    res = {}
    for c in cs:
        if c.observaciones:
            res[c.expediente] = "OBSERVADO (no generado)"
            continue
        ruta = out / f"R1 {c.expediente}-CC1-APE.docx"
        construir_docx(c, GUION[c.supuesto](c), RAIZ / "plantillas" / PLANTILLA[c.supuesto], ruta)
        err = verificar(ruta, c)
        if not a.sin_vista:
            err += vista(ruta, RAIZ / "plantillas" / PLANTILLA[c.supuesto], out / "_vista")
        res[c.expediente] = "GENERADO · APTO" if not err else "GENERADO · NO APTO: " + " / ".join(err)
        print(f"{c.expediente}  {c.supuesto}  {res[c.expediente]}  → {ruta.name}")
    (out / "_REPORTE.md").write_text(reporte(cs, fe, res), encoding="utf-8")
    print(f"Reporte: {out / '_REPORTE.md'}")


def vista(ruta, plantilla, carpeta):
    """Captura lado a lado frente a la plantilla y ficha de formato; devuelve diferencias de formato como errores."""
    import html as _h
    import vista as V
    carpeta.mkdir(parents=True, exist_ok=True)
    doc = (f"<!doctype html><meta charset='utf-8'><style>{V.CSS}</style><body>"
           + V.pagina_html(ruta, "GENERADA · " + ruta.name) + V.pagina_html(plantilla, "PLANTILLA · " + plantilla.name) + "</body>")
    (carpeta / f"{ruta.stem}.html").write_text(doc, encoding="utf-8")
    fic = V.ficha(ruta, plantilla)
    (carpeta / f"{ruta.stem}_FORMATO.md").write_text(fic, encoding="utf-8")
    try:
        aviso = V.capturar(doc, carpeta / f"{ruta.stem}.png")
        if aviso:
            print("   (vista) " + aviso)
    except Exception as e:  # la captura no bloquea: el HTML y la ficha ya están
        print(f"   (vista) sin captura: {type(e).__name__}: {e}")
    difs = [l[2:] for l in fic.split("## Diferencias", 1)[-1].splitlines() if l.startswith("- ")]
    # la plantilla trae OTRO caso: solo cuentan diferencias de tipografía/sangría, no de texto
    return [f"Formato distinto de la plantilla: {d_}" for d_ in difs if "tipografía" in d_ or "sangría" in d_ or "Página" in d_]


def cmd_verificar(a):
    malos = 0
    for r in a.docx:
        err = verificar(r)
        print(f"{'APTO' if not err else 'NO APTO'}  {r}")
        for e in err:
            print(f"    ✗ {e}")
        malos += bool(err)
    sys.exit(1 if malos else 0)


if __name__ == "__main__":
    A = args()
    {"datos": cmd_datos, "leer": cmd_leer, "generar": cmd_generar, "verificar": cmd_verificar}[A.accion](A)
